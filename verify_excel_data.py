#!/usr/bin/env python3
"""
验证Excel数据的完整性和格式
"""

import pandas as pd
import openpyxl
from openpyxl.styles import NamedStyle
import sys
import os

def verify_excel_data(excel_path):
    """验证Excel数据的完整性"""
    print(f"=== 验证Excel数据: {excel_path} ===")
    
    if not os.path.exists(excel_path):
        print(f"❌ 文件不存在: {excel_path}")
        return
    
    try:
        # 读取Excel文件
        df = pd.read_excel(excel_path, sheet_name='Document_Details')
        
        print(f"✅ 成功读取Excel文件")
        print(f"行数: {len(df)}")
        print(f"列数: {len(df.columns)}")
        
        # 重点验证第7-12行的数据（对应您截图中的红框区域）
        print(f"\n{'='*80}")
        print(f"重点验证第7-12行数据（红框区域）")
        print(f"{'='*80}")
        
        target_rows = [6, 7, 8, 9, 10, 11]  # 0-based索引
        
        for row_idx in target_rows:
            if row_idx < len(df):
                row = df.iloc[row_idx]
                
                std_amount = row.get('std_amount', 'N/A')
                pred_amount = row.get('pred_amount', 'N/A')
                check_amount = row.get('check_amount', 'N/A')
                
                std_trade_date = row.get('std_tradeDate', 'N/A')
                pred_trade_date = row.get('pred_tradeDate', 'N/A')
                check_trade_date = row.get('check_tradeDate', 'N/A')
                
                print(f"第{row_idx+1}行:")
                print(f"  amount: std={std_amount} pred={pred_amount} check={check_amount}")
                print(f"  tradeDate: std='{std_trade_date}' pred='{pred_trade_date}' check={check_trade_date}")
                
                # 验证数据逻辑
                if check_amount == False and pred_amount == 0:
                    print(f"  ❌ 问题: pred_amount不应该为0")
                elif check_amount == False and pred_amount != 0:
                    print(f"  ✅ 正常: pred_amount有实际值")
                elif check_amount == True:
                    print(f"  ✅ 正常: 匹配的数据")
        
        # 检查是否有任何pred_amount为0的异常情况
        print(f"\n{'='*80}")
        print(f"检查pred_amount为0的异常情况")
        print(f"{'='*80}")
        
        zero_amount_rows = []
        for i, row in df.iterrows():
            pred_amount = row.get('pred_amount', None)
            check_amount = row.get('check_amount', None)
            
            if pred_amount == 0 and check_amount == False:
                zero_amount_rows.append(i + 1)
        
        if zero_amount_rows:
            print(f"❌ 发现异常: 第{zero_amount_rows}行的pred_amount为0，但check_amount为False")
            print(f"这些行应该显示实际的预测金额，而不是0")
        else:
            print(f"✅ 没有发现pred_amount异常为0的情况")
        
        # 验证数据类型
        print(f"\n{'='*80}")
        print(f"验证数据类型")
        print(f"{'='*80}")
        
        amount_col = df['pred_amount']
        print(f"pred_amount列的数据类型: {amount_col.dtype}")
        print(f"pred_amount列的统计信息:")
        print(amount_col.describe())
        
        # 检查是否有NaN值
        nan_count = amount_col.isna().sum()
        print(f"pred_amount列的NaN值数量: {nan_count}")
        
        # 检查具体的值
        print(f"\npred_amount列的所有值:")
        for i, val in enumerate(amount_col):
            print(f"  第{i+1}行: {val} (类型: {type(val).__name__})")
        
        # 使用openpyxl直接读取原始数据
        print(f"\n{'='*80}")
        print(f"使用openpyxl直接读取原始数据")
        print(f"{'='*80}")
        
        workbook = openpyxl.load_workbook(excel_path)
        worksheet = workbook['Document_Details']
        
        # 找到pred_amount列的索引
        header_row = 1
        pred_amount_col = None
        for col in range(1, worksheet.max_column + 1):
            cell_value = worksheet.cell(row=header_row, column=col).value
            if cell_value == 'pred_amount':
                pred_amount_col = col
                break
        
        if pred_amount_col:
            print(f"pred_amount列在第{pred_amount_col}列")
            print(f"原始单元格数据:")
            for row in range(2, min(14, worksheet.max_row + 1)):  # 第2行到第13行
                cell = worksheet.cell(row=row, column=pred_amount_col)
                cell_value = cell.value
                cell_type = type(cell_value).__name__
                print(f"  第{row}行: {cell_value} (类型: {cell_type})")
        
        workbook.close()
        
    except Exception as e:
        print(f"❌ 验证过程中出错: {e}")
        import traceback
        traceback.print_exc()

if __name__ == "__main__":
    excel_path = "/var/folders/sn/shdg2z_d4ds1b5l83wjxpbvr0000gn/T/tmpuenwe4uw.xlsx"
    verify_excel_data(excel_path) 