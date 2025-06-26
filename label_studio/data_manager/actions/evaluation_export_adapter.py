"""
评估报告导出适配器
用于将ext_export_converter的功能集成到评估报告中
"""

import json
import tempfile
import os
from openpyxl import load_workbook
from label_studio.data_export.ext_export_converter import extract_annotations_and_data, fields


def create_label_studio_json_from_eval_data(eval_list, project_id=None):
    """
    将评估数据转换为Label Studio导出格式的JSON
    
    :param eval_list: Dictionary of filename -> (annotation_text, prediction_text, prompt_name)
    :param project_id: 项目ID
    :return: 符合Label Studio格式的数据列表
    """
    label_studio_data = []
    
    for idx, (filename, (annotation_text, prediction_text, prompt_name)) in enumerate(eval_list.items(), 1):
        # 构造Label Studio标准格式
        item = {
            "id": idx,
            "project": project_id,
            "data": {
                "filename": filename
            },
            "annotations": [
                {
                    "id": f"ann_{idx}",
                    "was_cancelled": False,
                    "result": [
                        {
                            "value": {
                                "text": [annotation_text]
                            }
                        }
                    ]
                }
            ],
            "predictions": [
                {
                    "id": f"pred_{idx}",
                    "result": [
                        {
                            "value": {
                                "text": [prediction_text]
                            }
                        }
                    ],
                    "prompt_name": prompt_name
                }
            ]
        }
        label_studio_data.append(item)
    
    return label_studio_data


def generate_annotations_predictions_sheets(eval_list, project_id=None):
    """
    基于评估数据生成Annotations和Predictions工作表
    
    :param eval_list: Dictionary of filename -> (annotation_text, prediction_text, prompt_name)
    :param project_id: 项目ID
    :return: tuple (annotations_worksheet, predictions_worksheet)
    """
    from openpyxl import Workbook
    
    # 转换数据格式
    label_studio_data = create_label_studio_json_from_eval_data(eval_list, project_id)
    
    # 创建临时JSON文件
    temp_fd, temp_json_path = tempfile.mkstemp(suffix='.json')
    try:
        with os.fdopen(temp_fd, 'w', encoding='utf-8') as f:
            json.dump(label_studio_data, f, ensure_ascii=False, indent=2)
        
        # 使用现有的extract_annotations_and_data功能
        results = extract_annotations_and_data(temp_json_path)
        
        # 创建工作簿和工作表
        wb = Workbook()
        ws_annotation = wb.active
        ws_annotation.title = "Annotations"
        ws_prediction = wb.create_sheet(title="Predictions")
        
        # 添加表头
        for ws in [ws_annotation, ws_prediction]:
            for col_idx, field in enumerate(fields, 1):
                ws.cell(row=1, column=col_idx, value=field)
        
        # 处理数据
        annotation_row = 2
        prediction_row = 2
        
        def process_text_data_internal(text, worksheet, start_row, project, item_id, filename):
            """内部数据处理函数"""
            current_row = start_row
            try:
                data = json.loads(text)
                if isinstance(data, list):
                    if len(data) == 0:
                        worksheet.cell(row=current_row, column=1, value=project)
                        worksheet.cell(row=current_row, column=2, value=item_id)
                        worksheet.cell(row=current_row, column=3, value=filename)
                        current_row += 1
                    else:
                        for item in data:
                            worksheet.cell(row=current_row, column=1, value=project)
                            worksheet.cell(row=current_row, column=2, value=item_id)
                            worksheet.cell(row=current_row, column=3, value=filename)
                            
                            if "page" not in item:
                                worksheet.cell(row=current_row, column=4, value=1)
                            
                            for col_idx, field in enumerate(fields[3:], 4):
                                if field in item:
                                    value = item[field]
                                    cell = worksheet.cell(row=current_row, column=col_idx)
                                    
                                    if isinstance(value, (int, float)):
                                        cell.value = value
                                        cell.number_format = '0.00'
                                    elif isinstance(value, dict) or isinstance(value, list):
                                        cell.value = json.dumps(value, ensure_ascii=False, indent=2, sort_keys=True)
                                    else:
                                        cell.value = str(value)
                            
                            current_row += 1
            except json.JSONDecodeError:
                worksheet.cell(row=current_row, column=1, value=project)
                worksheet.cell(row=current_row, column=2, value=item_id)
                worksheet.cell(row=current_row, column=3, value=filename)
                current_row += 1
            return current_row
        
        # 填充数据
        for project, item_id, filename, annotation_text, prediction_text in results:
            annotation_row = process_text_data_internal(
                annotation_text, ws_annotation, annotation_row, project, item_id, filename
            )
            prediction_row = process_text_data_internal(
                prediction_text, ws_prediction, prediction_row, project, item_id, filename
            )
        
        return ws_annotation, ws_prediction
        
    finally:
        # 清理临时文件
        if os.path.exists(temp_json_path):
            os.unlink(temp_json_path)


def add_annotations_predictions_to_excel(excel_path, eval_list, project_id=None):
    """
    将Annotations和Predictions工作表添加到现有的Excel文件中
    
    :param excel_path: 现有Excel文件路径
    :param eval_list: Dictionary of filename -> (annotation_text, prediction_text, prompt_name)
    :param project_id: 项目ID
    :return: 更新后的Excel文件路径
    """
    try:
        # 加载现有的Excel文件
        wb = load_workbook(excel_path)
        
        # 生成Annotations和Predictions工作表
        ws_annotation, ws_prediction = generate_annotations_predictions_sheets(eval_list, project_id)
        
        # 检查是否已存在同名工作表，如果存在则删除
        if 'Annotations' in wb.sheetnames:
            wb.remove(wb['Annotations'])
        if 'Predictions' in wb.sheetnames:
            wb.remove(wb['Predictions'])
        
        # 复制工作表到目标工作簿
        # 创建新的工作表
        target_ws_annotation = wb.create_sheet('Annotations')
        target_ws_prediction = wb.create_sheet('Predictions')
        
        # 复制Annotations工作表
        for row in ws_annotation.iter_rows():
            for cell in row:
                new_cell = target_ws_annotation.cell(row=cell.row, column=cell.column)
                new_cell.value = cell.value
                if cell.number_format:
                    new_cell.number_format = cell.number_format
        
        # 复制Predictions工作表
        for row in ws_prediction.iter_rows():
            for cell in row:
                new_cell = target_ws_prediction.cell(row=cell.row, column=cell.column)
                new_cell.value = cell.value
                if cell.number_format:
                    new_cell.number_format = cell.number_format
        
        # 保存文件
        wb.save(excel_path)
        
        return excel_path
        
    except Exception as e:
        print(f"添加Annotations和Predictions工作表时出错: {e}")
        return excel_path


def export_eval_data_to_excel(eval_list, output_path, project_id=None):
    """
    直接将评估数据导出为包含Annotations和Predictions的Excel文件
    
    :param eval_list: Dictionary of filename -> (annotation_text, prediction_text, prompt_name)
    :param output_path: 输出Excel文件路径
    :param project_id: 项目ID
    :return: Excel文件路径
    """
    try:
        # 生成工作表
        ws_annotation, ws_prediction = generate_annotations_predictions_sheets(eval_list, project_id)
        
        # 获取工作簿
        wb = ws_annotation.parent
        
        # 保存文件
        wb.save(output_path)
        
        return output_path
        
    except Exception as e:
        print(f"导出评估数据到Excel时出错: {e}")
        return None